#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_all_banks.py
====================
Standalone, generalised PDF bank statement extractor.
Supports Axis Bank, Bank of Baroda, HDFC Bank, ICICI Bank, Kotak Bank,
SBI (Layout 1 & 2), and Indian Bank statements.
Generates a styled Excel workbook with one sheet per PDF.
"""

import os
import re
import logging
from pathlib import Path
import pandas as pd
import pdfplumber
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ==============================================================================
# CONFIGURATION: Set this path to a directory containing PDF files,
# or directly to a single PDF bank statement file.
INPUT_PATH = "/Users/parshvapatel/Downloads/bank_statement_pdfs/bankofbaroda.pdf"
# ==============================================================================

# Set up logging
logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

# Canonical schema
CANONICAL_COLUMNS = [
    "DATE",
    "VALUE_DATE",
    "DESCRIPTION",
    "TRANSACTION_NO",
    "WITHDRAWAL_AMOUNT",
    "DEPOSIT_AMOUNT",
    "CLOSING_BALANCE",
]

# Date matching regexes
_DATE_PATTERNS = [
    re.compile(r"^\d{1,2}[-/\.]\d{1,2}[-/\.]\d{2,4}$"),  # 01/05/26, 01.05.2026, 01-10-2023
    re.compile(r"^\d{1,2}[-/][A-Za-z]{3}[-/]\d{2,4}$"),  # 02-Mar-2026
    re.compile(r"^\d{1,2}\s+[A-Za-z]{3}\s+\d{2,4}$"),    # 02 Mar 2026, 01 Jan 2026
    re.compile(r"^\d{4}[-/]\d{2}[-/]\d{2}$"),           # 2026-05-01
]

_AMOUNT_RE = re.compile(r"^[\d,]+\.?\d*$")
_CRDR_RE   = re.compile(r"\s*(cr|dr)$", re.IGNORECASE)
_DASH_RE   = re.compile(r"^[-–—]+$")

# Helper functions
def _collapse_whitespace(text: str) -> str:
    if not text:
        return ""
    return " ".join(text.split())

def _normalize_text(text: str) -> str:
    if not text:
        return ""
    text = text.lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()

def _is_date(text: str) -> bool:
    val = text.strip()
    return any(p.match(val) for p in _DATE_PATTERNS)

def _clean_amount(text: str) -> str:
    if not text:
        return ""
    val = _collapse_whitespace(text).replace(",", "")
    val = _CRDR_RE.sub("", val).strip()
    if _DASH_RE.match(val) or val == "" or val.lower() == "null":
        return ""
    return val

def _group_words_by_line(words: list[dict], y_tolerance: float = 3.0) -> list[dict]:
    lines: list[dict] = []
    for word in words:
        placed = False
        word_top = word["top"]
        for line in lines:
            if abs(line["top"] - word_top) <= y_tolerance:
                line["words"].append(word)
                placed = True
                break
        if not placed:
            lines.append({"top": word_top, "words": [word]})

    # Sort words within each line left-to-right
    for line in lines:
        line["words"] = sorted(line["words"], key=lambda w: w["x0"])

    # Compute line boundaries
    for line in lines:
        line["bottom"] = max(w["bottom"] for w in line["words"])
        line["top"] = min(w["top"] for w in line["words"])
        line["text"] = " ".join(w["text"] for w in line["words"])

    return sorted(lines, key=lambda l: l["top"])

def _classify_word(x0: float, column_ranges: dict[str, tuple[float, float]]) -> str | None:
    for col_name, (x_start, x_end) in column_ranges.items():
        if x_start <= x0 < x_end:
            return col_name
    return None

def _assign_narrations(transactions: list[dict], narration_buffer: list[dict], tolerance: float = 6.0):
    """Assign buffered narration words to their correct transaction using visual grid preceding."""
    transactions.sort(key=lambda t: t["center"])
    for nword in narration_buffer:
        y = nword["center"]
        target_txn = None
        for i, txn in enumerate(transactions):
            y_curr = txn["center"]
            y_next = transactions[i+1]["center"] if i + 1 < len(transactions) else float("inf")
            if (y_curr - tolerance) <= y < (y_next - tolerance):
                target_txn = txn
                break
        
        # Fallback to closest if slightly above the first transaction but close to it
        if target_txn is None and transactions:
            if y < transactions[0]["center"] - tolerance and (transactions[0]["center"] - y) <= 25.0:
                target_txn = transactions[0]
                
        if target_txn is not None:
            existing = target_txn["DESCRIPTION"]
            sep = " " if existing else ""
            target_txn["DESCRIPTION"] = existing + sep + nword["text"]

def detect_bank(pdf_path: str) -> str:
    try:
        with pdfplumber.open(pdf_path) as pdf:
            p = pdf.pages[0]
            text = (p.extract_text() or "").lower()
            header_text = text[:1000]
            
            scores = {}
            for bank_id, signatures in {
                "AXIS": ["axis bank", "axisbank", "utib000"],
                "BOB": ["bank of baroda", "barb0", "relationship type", "zqcd"],
                "HDFC": ["hdfc bank", "hdfcbank", "hdfc000"],
                "ICICI": ["icici bank", "icicibank", "saving account no. 165501501382"],
                "KOTAK": ["kotak mahindra", "kotak bank", "kkbk000"],
                "SBI": ["state bank of india", "welcome mr.", "welcome:", "sbi."],
                "INDIAN_BANK": ["indian bank", "idib000"],
            }.items():
                score = sum(1 for sig in signatures if sig in header_text)
                scores[bank_id] = score
                
            best_bank = max(scores, key=scores.get)
            if scores[best_bank] > 0:
                if best_bank == "SBI":
                    full_text = " ".join((page.extract_text() or "").lower() for page in pdf.pages[:3])
                    if "holding status" in full_text or "transaction details" in full_text:
                        return "SBI_LAYOUT_2"
                    return "SBI_LAYOUT_1"
                return best_bank
    except Exception as e:
        logger.error(f"Error detecting bank for {pdf_path}: {e}")
    return "UNKNOWN"

# --- AXIS BANK PROCESSOR ---
def process_axis(pdf_path: str) -> pd.DataFrame:
    logger.info(f"Processing {pdf_path} as Axis Bank statement...")
    all_rows = []
    
    col_mapping = {
        "DATE": ["tran date", "date"],
        "DESCRIPTION": ["particulars", "narration", "description"],
        "TRANSACTION_NO": ["chq no", "instr no", "ref no"],
        "WITHDRAWAL_AMOUNT": ["debit"],
        "DEPOSIT_AMOUNT": ["credit"],
        "CLOSING_BALANCE": ["balance"],
    }
    
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables() or []
            for table in tables:
                if not table or len(table[0]) < 4:
                    continue
                
                # Check if this table has header row
                header_row = table[0]
                resolved_cols = {}
                for idx, cell in enumerate(header_row):
                    val = _normalize_text(cell or "")
                    for canonical, aliases in col_mapping.items():
                        if any(alias in val for alias in aliases):
                            resolved_cols[canonical] = idx
                
                if len(resolved_cols) < 3:
                    if not all_rows:
                        continue
                
                start_idx = 1
                for row in table[start_idx:]:
                    if not row:
                        continue
                    
                    record = {col: "" for col in CANONICAL_COLUMNS}
                    for canonical, idx in resolved_cols.items():
                        if idx < len(row) and row[idx] is not None:
                            record[canonical] = _collapse_whitespace(str(row[idx]))
                    
                    desc = record["DESCRIPTION"].upper()
                    if any(s in desc for s in ["OPENING BALANCE", "CLOSING BALANCE", "TRANSACTION TOTAL"]):
                        continue
                    
                    # Clean fields
                    record["DATE"] = record["DATE"].strip()
                    record["WITHDRAWAL_AMOUNT"] = _clean_amount(record["WITHDRAWAL_AMOUNT"])
                    record["DEPOSIT_AMOUNT"] = _clean_amount(record["DEPOSIT_AMOUNT"])
                    record["CLOSING_BALANCE"] = _clean_amount(record["CLOSING_BALANCE"])
                    
                    if not record["DATE"] and not record["WITHDRAWAL_AMOUNT"] and not record["DEPOSIT_AMOUNT"]:
                        continue
                    
                    all_rows.append(record)
                    
    return pd.DataFrame(all_rows)

# --- BANK OF BARODA PROCESSOR ---
def process_bob(pdf_path: str) -> pd.DataFrame:
    logger.info(f"Processing {pdf_path} as Bank of Baroda statement...")
    column_ranges = {
        "DATE": (30.0, 76.0),
        "DESCRIPTION": (76.0, 245.0),
        "TRANSACTION_NO": (245.0, 305.0),
        "WITHDRAWAL_AMOUNT": (305.0, 395.0),
        "DEPOSIT_AMOUNT": (395.0, 490.0),
        "CLOSING_BALANCE": (490.0, 600.0),
    }
    
    all_transactions = []
    
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(x_tolerance=3, y_tolerance=3)
            if not words:
                continue
            
            lines = _group_words_by_line(words)
            table_started = False
            
            page_transactions = []
            page_narration_buffer = []
            
            for line in lines:
                text_norm = _normalize_text(line["text"])
                if "narration" in text_norm and "balance" in text_norm:
                    table_started = True
                    continue
                
                if not table_started:
                    continue
                
                line_words = line["words"]
                line_center_y = (line["top"] + line["bottom"]) / 2.0
                
                # Check if this line has a date
                date_word = None
                for w in line_words:
                    if column_ranges["DATE"][0] <= w["x0"] < column_ranges["DATE"][1]:
                        if _is_date(w["text"]):
                            date_word = w
                            break
                
                if date_word:
                    txn = {
                        "center": line_center_y,
                        "DATE": date_word["text"],
                        "VALUE_DATE": "",
                        "DESCRIPTION": "",
                        "TRANSACTION_NO": "",
                        "WITHDRAWAL_AMOUNT": "",
                        "DEPOSIT_AMOUNT": "",
                        "CLOSING_BALANCE": "",
                    }
                    
                    for w in line_words:
                        col = _classify_word(w["x0"], column_ranges)
                        if col == "DESCRIPTION":
                            page_narration_buffer.append({"text": w["text"], "center": line_center_y})
                        elif col == "TRANSACTION_NO":
                            txn["TRANSACTION_NO"] = w["text"]
                        elif col == "WITHDRAWAL_AMOUNT" and _clean_amount(w["text"]):
                            txn["WITHDRAWAL_AMOUNT"] = _clean_amount(w["text"])
                        elif col == "DEPOSIT_AMOUNT" and _clean_amount(w["text"]):
                            txn["DEPOSIT_AMOUNT"] = _clean_amount(w["text"])
                        elif col == "CLOSING_BALANCE" and _clean_amount(w["text"]):
                            txn["CLOSING_BALANCE"] = _clean_amount(w["text"])
                    
                    page_transactions.append(txn)
                else:
                    # Collect narration line
                    for w in line_words:
                        if _classify_word(w["x0"], column_ranges) == "DESCRIPTION":
                            page_narration_buffer.append({"text": w["text"], "center": line_center_y})
            
            # Proximity assign descriptions for this page
            _assign_narrations(page_transactions, page_narration_buffer, tolerance=6.0)
            all_transactions.extend(page_transactions)

    # Filter transactions
    clean_txns = []
    for t in all_transactions:
        t["DESCRIPTION"] = _collapse_whitespace(t["DESCRIPTION"])
        dl = t["DESCRIPTION"].lower()
        if any(s in dl for s in ["opening balance", "closing balance", "brought forward", "carried forward"]):
            continue
        if not t["WITHDRAWAL_AMOUNT"] and not t["DEPOSIT_AMOUNT"] and not t["CLOSING_BALANCE"]:
            continue
        t.pop("center", None)
        clean_txns.append(t)
        
    return pd.DataFrame(clean_txns)

# --- HDFC BANK PROCESSOR ---
def process_hdfc(pdf_path: str) -> pd.DataFrame:
    logger.info(f"Processing {pdf_path} as HDFC Bank statement...")
    column_ranges = {
        "DATE": (30.0, 70.0),
        "DESCRIPTION": (70.0, 275.0),
        "TRANSACTION_NO": (275.0, 360.0),
        "VALUE_DATE": (360.0, 400.0),
        "WITHDRAWAL_AMOUNT": (400.0, 485.0),
        "DEPOSIT_AMOUNT": (485.0, 560.0),
        "CLOSING_BALANCE": (560.0, 650.0),
    }
    
    all_transactions = []
    
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(x_tolerance=3, y_tolerance=3)
            if not words:
                continue
            
            lines = _group_words_by_line(words)
            table_started = False
            
            page_transactions = []
            page_narration_buffer = []
            
            for line in lines:
                text_norm = _normalize_text(line["text"])
                if "narration" in text_norm and "closing" in text_norm:
                    table_started = True
                    continue
                
                if not table_started:
                    continue
                
                date_text = None
                line_words = line["words"]
                line_center_y = (line["top"] + line["bottom"]) / 2.0
                
                for w in line_words:
                    if column_ranges["DATE"][0] <= w["x0"] < column_ranges["DATE"][1]:
                        if _is_date(w["text"]):
                            date_text = w["text"]
                            break
                
                if date_text:
                    txn = {
                        "center": line_center_y,
                        "DATE": date_text,
                        "VALUE_DATE": "",
                        "DESCRIPTION": "",
                        "TRANSACTION_NO": "",
                        "WITHDRAWAL_AMOUNT": "",
                        "DEPOSIT_AMOUNT": "",
                        "CLOSING_BALANCE": "",
                    }
                    
                    for w in line_words:
                        col = _classify_word(w["x0"], column_ranges)
                        if col == "DESCRIPTION":
                            page_narration_buffer.append({"text": w["text"], "center": line_center_y})
                        elif col == "TRANSACTION_NO":
                            txn["TRANSACTION_NO"] = w["text"]
                        elif col == "VALUE_DATE":
                            txn["VALUE_DATE"] = w["text"]
                        elif col == "WITHDRAWAL_AMOUNT" and _clean_amount(w["text"]):
                            txn["WITHDRAWAL_AMOUNT"] = _clean_amount(w["text"])
                        elif col == "DEPOSIT_AMOUNT" and _clean_amount(w["text"]):
                            txn["DEPOSIT_AMOUNT"] = _clean_amount(w["text"])
                        elif col == "CLOSING_BALANCE" and _clean_amount(w["text"]):
                            txn["CLOSING_BALANCE"] = _clean_amount(w["text"])
                            
                    page_transactions.append(txn)
                else:
                    for w in line_words:
                        if _classify_word(w["x0"], column_ranges) == "DESCRIPTION":
                            page_narration_buffer.append({"text": w["text"], "center": line_center_y})
            
            # Proximity assign descriptions for this page
            _assign_narrations(page_transactions, page_narration_buffer, tolerance=6.0)
            all_transactions.extend(page_transactions)

    clean_txns = []
    for t in all_transactions:
        t["DESCRIPTION"] = _collapse_whitespace(t["DESCRIPTION"])
        dl = t["DESCRIPTION"].lower()
        if any(s in dl for s in ["opening balance", "closing balance", "brought forward", "carried forward"]):
            continue
        if not t["WITHDRAWAL_AMOUNT"] and not t["DEPOSIT_AMOUNT"] and not t["CLOSING_BALANCE"]:
            continue
        t.pop("center", None)
        clean_txns.append(t)
        
    return pd.DataFrame(clean_txns)

# --- ICICI BANK PROCESSOR ---
def process_icici(pdf_path: str) -> pd.DataFrame:
    logger.info(f"Processing {pdf_path} as ICICI Bank statement...")
    column_ranges = {
        "DATE": (49.0, 119.0),
        "TRANSACTION_NO": (119.0, 189.0),
        "DESCRIPTION": (189.0, 390.0),
        "WITHDRAWAL_AMOUNT": (390.0, 456.0),
        "DEPOSIT_AMOUNT": (456.0, 522.0),
        "CLOSING_BALANCE": (522.0, 600.0),
    }
    
    all_transactions = []
    
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(x_tolerance=3, y_tolerance=3)
            if not words:
                continue
            
            lines = _group_words_by_line(words)
            table_started = False
            
            page_transactions = []
            page_narration_buffer = []
            
            for line in lines:
                text_norm = _normalize_text(line["text"])
                if "transaction remarks" in text_norm or ("remarks" in text_norm and "withdrawal" in text_norm):
                    table_started = True
                    continue
                
                if not table_started:
                    continue
                
                date_text = None
                line_words = line["words"]
                line_center_y = (line["top"] + line["bottom"]) / 2.0
                
                for w in line_words:
                    if column_ranges["DATE"][0] <= w["x0"] < column_ranges["DATE"][1]:
                        if _is_date(w["text"]) or re.match(r"^\d{2}\.\d{2}\.\d{4}$", w["text"]):
                            date_text = w["text"]
                            break
                            
                if date_text:
                    txn = {
                        "center": line_center_y,
                        "DATE": date_text,
                        "VALUE_DATE": "",
                        "DESCRIPTION": "",
                        "TRANSACTION_NO": "",
                        "WITHDRAWAL_AMOUNT": "",
                        "DEPOSIT_AMOUNT": "",
                        "CLOSING_BALANCE": "",
                    }
                    
                    for w in line_words:
                        col = _classify_word(w["x0"], column_ranges)
                        if col == "DESCRIPTION":
                            page_narration_buffer.append({"text": w["text"], "center": line_center_y})
                        elif col == "TRANSACTION_NO":
                            txn["TRANSACTION_NO"] = w["text"]
                        elif col == "WITHDRAWAL_AMOUNT" and _clean_amount(w["text"]):
                            txn["WITHDRAWAL_AMOUNT"] = _clean_amount(w["text"])
                        elif col == "DEPOSIT_AMOUNT" and _clean_amount(w["text"]):
                            txn["DEPOSIT_AMOUNT"] = _clean_amount(w["text"])
                        elif col == "CLOSING_BALANCE" and _clean_amount(w["text"]):
                            txn["CLOSING_BALANCE"] = _clean_amount(w["text"])
                            
                    page_transactions.append(txn)
                else:
                    for w in line_words:
                        if _classify_word(w["x0"], column_ranges) == "DESCRIPTION":
                            page_narration_buffer.append({"text": w["text"], "center": line_center_y})
            
            # Proximity assign descriptions for this page
            _assign_narrations(page_transactions, page_narration_buffer, tolerance=6.0)
            all_transactions.extend(page_transactions)

    clean_txns = []
    for t in all_transactions:
        t["DESCRIPTION"] = _collapse_whitespace(t["DESCRIPTION"])
        dl = t["DESCRIPTION"].lower()
        if any(s in dl for s in ["opening balance", "closing balance", "brought forward", "carried forward"]):
            continue
        if not t["WITHDRAWAL_AMOUNT"] and not t["DEPOSIT_AMOUNT"] and not t["CLOSING_BALANCE"]:
            continue
        t.pop("center", None)
        clean_txns.append(t)
        
    return pd.DataFrame(clean_txns)

# --- KOTAK BANK PROCESSOR ---
def process_kotak(pdf_path: str) -> pd.DataFrame:
    logger.info(f"Processing {pdf_path} as Kotak Bank statement...")
    all_rows = []
    
    col_mapping = {
        "DATE": ["date"],
        "DESCRIPTION": ["description", "particulars", "narration"],
        "TRANSACTION_NO": ["chq/ref. no.", "chq/ref no", "ref no"],
        "WITHDRAWAL_AMOUNT": ["withdrawal (dr.)", "withdrawal dr", "withdrawal"],
        "DEPOSIT_AMOUNT": ["deposit (cr.)", "deposit cr", "deposit"],
        "CLOSING_BALANCE": ["balance"],
    }
    
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables() or []
            for table in tables:
                if not table or len(table[0]) < 4:
                    continue
                
                header_row = None
                resolved_cols = {}
                header_idx = -1
                
                for idx, row in enumerate(table[:3]):
                    row_text = " ".join(_normalize_text(cell or "") for cell in row)
                    if "date" in row_text and "balance" in row_text:
                        header_row = row
                        header_idx = idx
                        break
                
                if header_row is None:
                    if len(table[0]) == 7:
                        resolved_cols = {
                            "DATE": 1,
                            "DESCRIPTION": 2,
                            "TRANSACTION_NO": 3,
                            "WITHDRAWAL_AMOUNT": 4,
                            "DEPOSIT_AMOUNT": 5,
                            "CLOSING_BALANCE": 6,
                        }
                        header_idx = 0 if "Savings" in str(table[0][0]) else -1
                    else:
                        continue
                else:
                    for idx, cell in enumerate(header_row):
                        val = _normalize_text(cell or "")
                        for canonical, aliases in col_mapping.items():
                            if any(alias in val for alias in aliases):
                                resolved_cols[canonical] = idx
                
                start_idx = header_idx + 1
                for row in table[start_idx:]:
                    if not row or len(row) < 7:
                        continue
                    
                    serial = (row[0] or "").strip()
                    if serial in ("", "-") and all_rows:
                        desc_idx = resolved_cols.get("DESCRIPTION", 2)
                        cont_text = (row[desc_idx] or "").strip()
                        if cont_text:
                            all_rows[-1]["DESCRIPTION"] += " " + cont_text
                        continue
                        
                    record = {col: "" for col in CANONICAL_COLUMNS}
                    for canonical, idx in resolved_cols.items():
                        if idx < len(row) and row[idx] is not None:
                            record[canonical] = _collapse_whitespace(str(row[idx]))
                    
                    record["DATE"] = _collapse_whitespace(record["DATE"])
                    if not _is_date(record["DATE"]):
                        continue
                        
                    record["WITHDRAWAL_AMOUNT"] = _clean_amount(record["WITHDRAWAL_AMOUNT"])
                    record["DEPOSIT_AMOUNT"] = _clean_amount(record["DEPOSIT_AMOUNT"])
                    record["CLOSING_BALANCE"] = _clean_amount(record["CLOSING_BALANCE"])
                    
                    if not record["WITHDRAWAL_AMOUNT"] and not record["DEPOSIT_AMOUNT"] and not record["CLOSING_BALANCE"]:
                        continue
                        
                    all_rows.append(record)
                    
    clean_rows = []
    for r in all_rows:
        dl = r["DESCRIPTION"].lower()
        if any(s in dl for s in ["opening balance", "closing balance"]):
            continue
        clean_rows.append(r)
        
    return pd.DataFrame(clean_rows)

# --- SBI LAYOUT 1 PROCESSOR (IMAGE HEADER) ---
def process_sbi_layout_1(pdf_path: str) -> pd.DataFrame:
    logger.info(f"Processing {pdf_path} as SBI Layout 1 statement...")
    all_rows = []
    
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables() or []
            for table in tables:
                if not table or len(table[0]) < 7:
                    continue
                
                is_sbi_table = False
                first_row = table[0]
                last_cell = _normalize_text(first_row[-1] or "")
                if last_cell == "balance" or (first_row[0] == "" and first_row[1] == "" and last_cell == "balance"):
                    is_sbi_table = True
                
                if not is_sbi_table:
                    for row in table[:2]:
                        if _is_date(str(row[0])):
                            is_sbi_table = True
                            break
                            
                if not is_sbi_table:
                    continue
                
                start_row_idx = 1 if (last_cell == "balance" or first_row[0] == "") else 0
                
                for row in table[start_row_idx:]:
                    if not row or len(row) < 7:
                        continue
                        
                    date_val = _collapse_whitespace(str(row[0] or ""))
                    if not _is_date(date_val):
                        continue
                        
                    record = {
                        "DATE": date_val,
                        "VALUE_DATE": _collapse_whitespace(str(row[1] or "")),
                        "DESCRIPTION": _collapse_whitespace(str(row[2] or "")),
                        "TRANSACTION_NO": _collapse_whitespace(str(row[3] or "")),
                        "WITHDRAWAL_AMOUNT": _clean_amount(str(row[4] or "")),
                        "DEPOSIT_AMOUNT": _clean_amount(str(row[5] or "")),
                        "CLOSING_BALANCE": _clean_amount(str(row[6] or "")),
                    }
                    
                    desc_lower = record["DESCRIPTION"].lower()
                    if any(s in desc_lower for s in ["brought forward", "carried forward", "statement summary"]):
                        continue
                        
                    if not record["WITHDRAWAL_AMOUNT"] and not record["DEPOSIT_AMOUNT"] and not record["CLOSING_BALANCE"]:
                        continue
                        
                    all_rows.append(record)
                    
    return pd.DataFrame(all_rows)

# --- SBI LAYOUT 2 PROCESSOR ---
def process_sbi_layout_2(pdf_path: str) -> pd.DataFrame:
    logger.info(f"Processing {pdf_path} as SBI Layout 2 statement...")
    all_rows = []
    
    col_mapping = {
        "DATE": ["date"],
        "DESCRIPTION": ["transaction reference", "particulars", "description"],
        "TRANSACTION_NO": ["ref.no./chq.no.", "ref no", "chq no"],
        "WITHDRAWAL_AMOUNT": ["debit"],
        "DEPOSIT_AMOUNT": ["credit"],
        "CLOSING_BALANCE": ["balance"],
    }
    
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables() or []
            for table in tables:
                if not table or len(table[0]) < 6:
                    continue
                
                header_row = None
                resolved_cols = {}
                header_idx = -1
                
                for idx, row in enumerate(table[:3]):
                    row_text = " ".join(_normalize_text(cell or "") for cell in row)
                    if "transaction reference" in row_text and "balance" in row_text:
                        header_row = row
                        header_idx = idx
                        break
                        
                if header_row is None:
                    if all_rows and len(table[0]) == 7:
                        resolved_cols = {
                            "DATE": 0,
                            "DESCRIPTION": 1,
                            "TRANSACTION_NO": 3,
                            "WITHDRAWAL_AMOUNT": 5,
                            "DEPOSIT_AMOUNT": 4,
                            "CLOSING_BALANCE": 6,
                        }
                        header_idx = -1
                    else:
                        continue
                else:
                    for idx, cell in enumerate(header_row):
                        val = _normalize_text(cell or "")
                        for canonical, aliases in col_mapping.items():
                            if any(alias in val for alias in aliases):
                                resolved_cols[canonical] = idx
                
                start_idx = header_idx + 1
                for row in table[start_idx:]:
                    if not row or len(row) < 5:
                        continue
                    
                    date_val = _collapse_whitespace(str(row[resolved_cols.get("DATE", 0)] or ""))
                    if not _is_date(date_val):
                        desc_idx = resolved_cols.get("DESCRIPTION", 1)
                        if desc_idx < len(row) and row[desc_idx] and all_rows:
                            cont_text = _collapse_whitespace(str(row[desc_idx]))
                            if cont_text and not cont_text.lower().startswith("opening balance"):
                                all_rows[-1]["DESCRIPTION"] += " " + cont_text
                        continue
                    
                    record = {col: "" for col in CANONICAL_COLUMNS}
                    for canonical, idx in resolved_cols.items():
                        if idx < len(row) and row[idx] is not None:
                            record[canonical] = _collapse_whitespace(str(row[idx]))
                    
                    record["WITHDRAWAL_AMOUNT"] = _clean_amount(record["WITHDRAWAL_AMOUNT"])
                    record["DEPOSIT_AMOUNT"] = _clean_amount(record["DEPOSIT_AMOUNT"])
                    record["CLOSING_BALANCE"] = _clean_amount(record["CLOSING_BALANCE"])
                    
                    desc_lower = record["DESCRIPTION"].lower()
                    if any(s in desc_lower for s in ["opening balance", "closing balance", "brought forward"]):
                        continue
                        
                    if not record["WITHDRAWAL_AMOUNT"] and not record["DEPOSIT_AMOUNT"] and not record["CLOSING_BALANCE"]:
                        continue
                        
                    all_rows.append(record)
                    
    return pd.DataFrame(all_rows)

# --- INDIAN BANK PROCESSOR ---
def process_indian_bank(pdf_path: str) -> pd.DataFrame:
    logger.info(f"Processing {pdf_path} as Indian Bank statement...")
    
    # Adjusted column boundaries to fix classification
    column_ranges = {
        "DATE": (36.0, 85.0),
        "VALUE_DATE": (85.0, 138.0),
        "DESCRIPTION": (138.0, 320.0),
        "TRANSACTION_NO": (320.0, 365.0),
        "WITHDRAWAL_AMOUNT": (365.0, 425.0),
        "DEPOSIT_AMOUNT": (425.0, 490.0),
        "CLOSING_BALANCE": (490.0, 565.0),
    }
    
    all_transactions = []
    
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(x_tolerance=3, y_tolerance=3)
            if not words:
                continue
            
            lines = _group_words_by_line(words)
            table_started = False
            
            page_transactions = []
            page_narration_buffer = []
            
            for line in lines:
                text_norm = _normalize_text(line["text"])
                if "post date" in text_norm and "balance" in text_norm:
                    table_started = True
                    continue
                
                if not table_started:
                    continue
                
                line_words = line["words"]
                line_center_y = (line["top"] + line["bottom"]) / 2.0
                
                post_date_text = None
                for w in line_words:
                    if column_ranges["DATE"][0] <= w["x0"] < column_ranges["DATE"][1]:
                        if _is_date(w["text"]):
                            post_date_text = w["text"]
                            break
                            
                if post_date_text:
                    txn = {
                        "center": line_center_y,
                        "DATE": post_date_text,
                        "VALUE_DATE": "",
                        "DESCRIPTION": "",
                        "TRANSACTION_NO": "",
                        "WITHDRAWAL_AMOUNT": "",
                        "DEPOSIT_AMOUNT": "",
                        "CLOSING_BALANCE": "",
                    }
                    
                    for w in line_words:
                        col = _classify_word(w["x0"], column_ranges)
                        if col == "DESCRIPTION":
                            page_narration_buffer.append({"text": w["text"], "center": line_center_y})
                        elif col == "TRANSACTION_NO":
                            txn["TRANSACTION_NO"] = w["text"]
                        elif col == "VALUE_DATE":
                            txn["VALUE_DATE"] = w["text"]
                        elif col == "WITHDRAWAL_AMOUNT" and _clean_amount(w["text"]):
                            txn["WITHDRAWAL_AMOUNT"] = _clean_amount(w["text"])
                        elif col == "DEPOSIT_AMOUNT" and _clean_amount(w["text"]):
                            txn["DEPOSIT_AMOUNT"] = _clean_amount(w["text"])
                        elif col == "CLOSING_BALANCE" and _clean_amount(w["text"]):
                            txn["CLOSING_BALANCE"] = _clean_amount(w["text"])
                            
                    page_transactions.append(txn)
                else:
                    for w in line_words:
                        if _classify_word(w["x0"], column_ranges) == "DESCRIPTION":
                            page_narration_buffer.append({"text": w["text"], "center": line_center_y})
            
            # Proximity assign descriptions for this page
            _assign_narrations(page_transactions, page_narration_buffer, tolerance=6.0)
            all_transactions.extend(page_transactions)

    clean_txns = []
    for t in all_transactions:
        t["DESCRIPTION"] = _collapse_whitespace(t["DESCRIPTION"])
        dl = t["DESCRIPTION"].lower()
        if any(s in dl for s in ["opening balance", "closing balance", "brought forward", "carried forward"]):
            continue
        if not t["WITHDRAWAL_AMOUNT"] and not t["DEPOSIT_AMOUNT"] and not t["CLOSING_BALANCE"]:
            continue
        t.pop("center", None)
        clean_txns.append(t)
        
    return pd.DataFrame(clean_txns)

# --- EXCEL WRITER AND STYLER ---
def write_to_excel(dfs_dict: dict[str, pd.DataFrame], output_path: str):
    logger.info(f"Writing all sheets to {output_path}...")
    
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        for sheet_name, df in dfs_dict.items():
            for col in CANONICAL_COLUMNS:
                if col not in df.columns:
                    df[col] = ""
            
            df = df[CANONICAL_COLUMNS]
            
            df_excel = df.copy()
            for col in ["WITHDRAWAL_AMOUNT", "DEPOSIT_AMOUNT", "CLOSING_BALANCE"]:
                df_excel[col] = pd.to_numeric(df_excel[col], errors='coerce')
            
            df_excel.to_excel(writer, sheet_name=sheet_name, index=False)
            
            workbook = writer.book
            worksheet = workbook[sheet_name]
            
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            for col_idx in range(1, len(CANONICAL_COLUMNS) + 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
            
            data_font = Font(name="Segoe UI", size=10)
            thin_border = Border(
                left=Side(style='thin', color='D3D3D3'),
                right=Side(style='thin', color='D3D3D3'),
                top=Side(style='thin', color='D3D3D3'),
                bottom=Side(style='thin', color='D3D3D3')
            )
            
            for r_idx in range(2, worksheet.max_row + 1):
                for c_idx in [1, 2, 3, 4]:
                    cell = worksheet.cell(row=r_idx, column=c_idx)
                    cell.font = data_font
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                
                for c_idx in [5, 6, 7]:
                    cell = worksheet.cell(row=r_idx, column=c_idx)
                    cell.font = data_font
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if cell.value is not None:
                        cell.number_format = '#,##0.00'
            
            for col in worksheet.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    val = str(cell.value or '')
                    if len(val) > max_len:
                        max_len = len(val)
                worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
                
            worksheet.freeze_panes = "A2"

# --- MAIN ORCHESTRATOR ---
def main():
    path = Path(INPUT_PATH)
    if not path.exists():
        logger.error(f"Provided path does not exist: {INPUT_PATH}")
        return

    if path.is_file():
        pdf_files = [path]
        output_xlsx = path.parent / "bank_statements_all.xlsx"
        logger.info(f"Processing single file: {path.name}")
    else:
        pdf_files = list(path.glob("*.pdf"))
        output_xlsx = path / "bank_statements_all.xlsx"
        logger.info(f"Found {len(pdf_files)} PDF files to process in directory.")
    
    results = {}
    for pdf_path in pdf_files:
        filename = pdf_path.name
        
        bank = detect_bank(str(pdf_path))
        logger.info(f"Detected bank for '{filename}': {bank}")
        
        df = pd.DataFrame()
        try:
            if bank == "AXIS":
                df = process_axis(str(pdf_path))
            elif bank == "BOB":
                df = process_bob(str(pdf_path))
            elif bank == "HDFC":
                df = process_hdfc(str(pdf_path))
            elif bank == "ICICI":
                df = process_icici(str(pdf_path))
            elif bank == "KOTAK":
                df = process_kotak(str(pdf_path))
            elif bank == "SBI_LAYOUT_1":
                df = process_sbi_layout_1(str(pdf_path))
            elif bank == "SBI_LAYOUT_2":
                df = process_sbi_layout_2(str(pdf_path))
            elif bank == "INDIAN_BANK":
                df = process_indian_bank(str(pdf_path))
            else:
                logger.warning(f"Unknown bank format for {filename}, skipping.")
                continue
                
            if not df.empty:
                sheet_name = pdf_path.stem[:30]
                results[sheet_name] = df
                logger.info(f"Successfully extracted {len(df)} rows from {filename}.")
            else:
                logger.warning(f"No transactions extracted from {filename}.")
        except Exception as e:
            logger.error(f"Failed to process {filename}: {e}", exc_info=True)
            
    if results:
        write_to_excel(results, str(output_xlsx))
        logger.info("Extraction complete! Output saved to: bank_statements_all.xlsx")
    else:
        logger.error("No data extracted from any PDF.")

if __name__ == "__main__":
    main()
